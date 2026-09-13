#!/usr/bin/env python3
"""
Validate the tumour-minus-normal protein bridge against an mRNA-protein anchor
and aggregate the TCGA 450K methylation layer to gene level.

Purpose
    The delta transformation is only usable as a cross-cohort bridge if it
    preserves real biology. This module tests that: within each CPTAC cohort it
    correlates the tumour-minus-normal protein contrast with the corresponding
    mRNA contrast, and compares the result with the raw protein-versus-mRNA
    correlation of the same cohort. A delta-level anchor correlation clearly
    above the raw one means the contrast recovers the transcript-protein
    agreement that cohort-specific references destroy. In the same run the
    module aggregates the TCGA-UCEC Illumina 450K beta matrix from probes to
    genes, using the probe-to-gene map of the pan-cancer assets, so that the
    methylation layer becomes comparable with the CPTAC gene-level methylation
    matrices; the cross-cohort methylation correlations close the validation.

Author
    Tao Zhu

Created
    2026-09-13

Inputs
    data("cptac_discovery_prot_tumor")     .cct, Discovery cohort tumour protein
    data("cptac_discovery_prot_normal")    .cct, Discovery cohort normal protein
    data("cptac_discovery_rna_tumor")      .cct, Discovery cohort tumour RNA-seq
    data("cptac_discovery_rna_normal")     .cct, Discovery cohort normal RNA-seq
    data("cptac_ov_prot_tumor")            .cct, Ovarian cohort tumour protein
    data("cptac_ov_prot_normal")           .cct, Ovarian cohort normal protein
    data("cptac_ov_rna")                   .cct, Ovarian cohort RNA-seq (no normal
        RNA-seq exists for this cohort, so only the raw comparison is reported)
    data("cptac_independent_prot_ratio")   .cct, Independent tumour/normal ratio
    data("cptac_independent_rna")          .cct, Independent tumour/normal ratio
    data("cptac_independent_metadata")     .xlsx, sample groups for the two above
    data("cptac_discovery_methylation")    .cct, Discovery gene-level beta values
    data("cptac_independent_methylation")  .cct, Independent gene-level beta values
    data("xena_methylation_probe_map")     text, probe-to-gene map (tab separated)
    data("xena_ucec_methylation_450k")     .gz, TCGA-UCEC 450K probe-level beta
        matrix, read in chunks

Outputs
    work("protein_bridge_anchors")         .json, anchor and methylation records
    work("tcga_methylation_genelevel")     .csv.gz, TCGA-UCEC gene-level beta
        matrix produced from the probe-level matrix

Usage
    python 02_03_validate_mrna_protein_anchor.py
"""

from __future__ import annotations

import gzip
import json
import re
import sys
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
from scipy.stats import spearmanr

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from common.config import data, ensure_dir, param, work  # noqa: E402

# Identifier vocabulary of the LinkedOmics gene-level matrices; see 02_01.
IDENTIFIER_PATTERN = re.compile(r"^[A-Za-z][A-Za-z0-9\-\._@]*$")

# Column names of the Independent-cohort metadata workbook and the group values
# that identify a tumour and a normal specimen; see 02_02.
METADATA_CASE_ID_COLUMN = "Case_id"
METADATA_GROUP_COLUMN = "Group"
TUMOUR_GROUP = "Tumor"
NORMAL_GROUPS = ("Adjacent_normal", "Enriched_Normal")
SAMPLE_SUFFIX_PATTERN = re.compile(r"-[A-Z]$")

# The TCGA 450K beta matrix holds far more probes than the gene space needs. It
# is read in chunks of this many rows and only probes present in the map are
# retained, so the full probe-level matrix is never held in memory at once.
METHYLATION_CHUNK_SIZE = 40000


def load_matrix(path: Path, columns: list[int] | None = None) -> pd.DataFrame:
    """Read a tab-separated gene-level matrix of a LinkedOmics export.

    Args:
        path: Matrix file; gzip compression is detected from the file suffix.
        columns: Optional list of column positions to read instead of all of
            them.

    Returns:
        DataFrame of float values indexed by gene symbol with one column per
        sample. Duplicate gene symbols are averaged and non-numeric entries
        become NaN.
    """
    frame = pd.read_csv(
        path,
        sep="\t",
        comment="#",
        header=0,
        dtype=str,
        engine="python",
        usecols=columns,
    )
    frame = frame.rename(columns={frame.columns[0]: "ID"})
    frame["ID"] = frame["ID"].astype(str).str.strip().str.strip('"')
    frame = frame[frame["ID"].str.match(IDENTIFIER_PATTERN, na=False)]
    matrix = frame.set_index("ID").apply(pd.to_numeric, errors="coerce")
    return matrix.groupby(level=0).mean()


def read_group_map(metadata_path: Path) -> dict[str, str]:
    """Read the case-identifier-to-group map of the Independent cohort.

    Args:
        metadata_path: Excel workbook with Case_id and Group columns.

    Returns:
        Mapping of upper-cased case identifier to its group label.
    """
    workbook = openpyxl.load_workbook(metadata_path, read_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    header = [str(value) for value in rows[0]]
    group_index = header.index(METADATA_GROUP_COLUMN)
    case_index = header.index(METADATA_CASE_ID_COLUMN)
    return {
        str(row[case_index]).strip().upper(): str(row[group_index])
        for row in rows[1:]
        if row[case_index]
    }


def split_by_group(matrix: pd.DataFrame, group_map: dict[str, str]) -> tuple:
    """Split a cohort matrix into tumour-mean and normal-mean spectra.

    Args:
        matrix: Gene-by-sample matrix of one cohort.
        group_map: Case identifier to group label, from read_group_map.

    Returns:
        Pair of gene-indexed Series: the mean over the tumour columns and the mean
        over the normal columns.
    """
    column_groups: dict[str, str] = {}
    for column in matrix.columns:
        # Sample identifiers carry a replicate suffix that the metadata does not.
        key = SAMPLE_SUFFIX_PATTERN.sub("", str(column).upper())
        if key in group_map:
            column_groups[column] = group_map[key]
    tumour_columns = [
        column for column, group in column_groups.items() if group == TUMOUR_GROUP
    ]
    normal_columns = [
        column for column, group in column_groups.items() if group in NORMAL_GROUPS
    ]
    return (
        matrix[tumour_columns].mean(axis=1, skipna=True),
        matrix[normal_columns].mean(axis=1, skipna=True),
    )


def correlate(a: pd.Series, b: pd.Series, label: str) -> dict | None:
    """Spearman-correlate two gene-indexed profiles on their shared genes.

    Args:
        a: First gene-indexed profile.
        b: Second gene-indexed profile.
        label: Text used in the console line.

    Returns:
        Record with the number of genes compared, the Spearman rho and its
        p-value, or None when the shared gene set is smaller than the minimum
        required by the pre-registration.
    """
    shared = a.index.intersection(b.index)
    x = a.reindex(shared).values.astype(float)
    y = b.reindex(shared).values.astype(float)
    finite = np.isfinite(x) & np.isfinite(y)
    if finite.sum() < param("protein_bridge", "min_shared_genes"):
        print("   %-30s too few shared genes (%d)" % (label, finite.sum()))
        return None
    rho, p_value = spearmanr(x[finite], y[finite])
    print("   %-30s n=%-6d rho=%6.3f  p=%.1e" % (label, finite.sum(), rho, p_value))
    return {"n": int(finite.sum()), "rho": float(rho), "p": float(p_value)}


def read_probe_to_gene(map_path: Path) -> dict[str, list[str]]:
    """Read the Illumina 450K probe-to-gene map.

    Args:
        map_path: Tab-separated text file whose first field is the probe
            identifier and second field a comma-separated gene list.

    Returns:
        Mapping of probe identifier to its non-empty gene list. Probes without a
        gene symbol (empty or ".") are skipped.
    """
    probe_to_genes: dict[str, list[str]] = {}
    with open(map_path, errors="replace") as handle:
        for line in handle:
            fields = line.rstrip("\n").split("\t")
            if len(fields) >= 2 and fields[1].strip() not in ("", "."):
                probe_to_genes[fields[0].strip()] = [
                    gene for gene in fields[1].split(",") if gene
                ]
    print(
        "   probe map: %d probes -> %d genes"
        % (
            len(probe_to_genes),
            len({gene for genes in probe_to_genes.values() for gene in genes}),
        )
    )
    return probe_to_genes


def aggregate_probes_to_genes(
    beta_path: Path, probe_to_genes: dict[str, list[str]]
) -> pd.DataFrame:
    """Aggregate a probe-level beta matrix to a gene-level beta matrix.

    Args:
        beta_path: Probe-by-sample beta matrix, tab separated, possibly with a
            leading comment block.
        probe_to_genes: Probe-to-gene map, from read_probe_to_gene.

    Returns:
        Gene-by-sample DataFrame. A probe annotating several genes contributes to
        its first gene only, and several probes of a gene are averaged.
    """
    kept_chunks: list[pd.DataFrame] = []
    n_read = 0
    n_kept = 0
    for chunk in pd.read_csv(
        beta_path,
        sep="\t",
        comment="#",
        header=0,
        dtype=str,
        engine="python",
        chunksize=METHYLATION_CHUNK_SIZE,
    ):
        n_read += len(chunk)
        probe_ids = chunk[chunk.columns[0]].astype(str).str.strip()
        kept = chunk[probe_ids.isin(probe_to_genes.keys())]
        n_kept += len(kept)
        if len(kept):
            kept_chunks.append(kept)
    print(
        "   rows scanned=%d  matched by the map=%d (%.1f%%)"
        % (n_read, n_kept, 100 * n_kept / max(n_read, 1))
    )
    probes = pd.concat(kept_chunks)
    probes = probes.rename(columns={probes.columns[0]: "probe"})
    numeric = probes.set_index("probe").apply(pd.to_numeric, errors="coerce")
    # Collect one value vector per gene; a multi-gene probe is assigned to its
    # first gene only, which keeps the gene assignment unambiguous.
    gene_vectors: dict[str, list[np.ndarray]] = {}
    for probe, row in numeric.iterrows():
        for gene in probe_to_genes.get(probe, [])[:1]:
            gene_vectors.setdefault(gene, []).append(row.values)
    gene_matrix = pd.DataFrame(
        {
            gene: np.nanmean(np.vstack(vectors), axis=0)
            for gene, vectors in gene_vectors.items()
        }
    ).T
    gene_matrix.columns = numeric.columns
    return gene_matrix


def main() -> None:
    """Run the mRNA-protein anchor validation, aggregate the TCGA 450K layer to
    gene level and write the anchor record file."""
    print("=" * 96)
    print("[E] Within-cohort mRNA-protein anchor (delta(T-N) level)")
    print("=" * 96)
    records: dict[str, dict | None] = {}

    # Discovery cohort: the delta contrast is available on both layers.
    dis_protein = load_matrix(data("cptac_discovery_prot_tumor")).mean(
        axis=1, skipna=True
    )
    dis_protein_normal = load_matrix(data("cptac_discovery_prot_normal")).mean(
        axis=1, skipna=True
    )
    dis_rna = load_matrix(data("cptac_discovery_rna_tumor")).mean(axis=1, skipna=True)
    dis_rna_normal = load_matrix(data("cptac_discovery_rna_normal")).mean(
        axis=1, skipna=True
    )
    records["dis_dProt_dRNA"] = correlate(
        dis_protein - dis_protein_normal,
        dis_rna - dis_rna_normal,
        "dis: delta protein vs delta mRNA",
    )
    records["dis_raw_Prot_RNA"] = correlate(
        dis_protein, dis_rna, "dis: raw protein vs raw mRNA"
    )

    # Ovarian cohort: no normal RNA-seq exists, so only the raw comparison can be
    # formed.
    ov_protein = load_matrix(data("cptac_ov_prot_tumor")).mean(axis=1, skipna=True)
    ov_protein_normal = load_matrix(data("cptac_ov_prot_normal")).mean(
        axis=1, skipna=True
    )
    ov_rna = load_matrix(data("cptac_ov_rna")).mean(axis=1, skipna=True)
    records["ov_raw_Prot_RNA"] = correlate(
        ov_protein, ov_rna, "ov: raw protein vs raw mRNA (no normal RNA)"
    )

    # Independent cohort: both matrices are tumour/normal ratios and are split by
    # the metadata groups.
    group_map = read_group_map(data("cptac_independent_metadata"))
    ind_protein_t, ind_protein_n = split_by_group(
        load_matrix(data("cptac_independent_prot_ratio")), group_map
    )
    ind_rna_t, ind_rna_n = split_by_group(
        load_matrix(data("cptac_independent_rna")), group_map
    )
    records["ind_dProt_dRNA"] = correlate(
        ind_protein_t - ind_protein_n,
        ind_rna_t - ind_rna_n,
        "ind: delta protein vs delta mRNA",
    )
    records["ind_raw_Prot_RNA"] = correlate(
        ind_protein_t, ind_rna_t, "ind: raw protein vs raw mRNA"
    )

    print("\n" + "=" * 96)
    print("[F] TCGA 450K probe-to-gene aggregation")
    print("=" * 96)
    probe_to_genes = read_probe_to_gene(data("xena_methylation_probe_map"))
    gene_matrix = aggregate_probes_to_genes(
        data("xena_ucec_methylation_450k"), probe_to_genes
    )
    print("   TCGA-UCEC gene-level methylation: %d genes x %d samples" % gene_matrix.shape)
    gene_means = gene_matrix.mean(axis=1, skipna=True)
    print(
        "   gene means range[%.3f, %.3f] median=%.3f"
        % (gene_means.min(), gene_means.max(), gene_means.median())
    )
    gene_level_path = work("tcga_methylation_genelevel")
    ensure_dir(gene_level_path.parent)
    gene_matrix.to_csv(gene_level_path, compression="gzip")

    print("\n" + "=" * 96)
    print("[G] Cross-cohort methylation agreement (including TCGA)")
    print("=" * 96)
    ind_methylation = load_matrix(data("cptac_independent_methylation")).mean(
        axis=1, skipna=True
    )
    dis_methylation = load_matrix(data("cptac_discovery_methylation")).mean(
        axis=1, skipna=True
    )
    records["meth_TCGA_ind"] = correlate(
        gene_means, ind_methylation, "TCGA450K(gene level) vs CPTAC-ind"
    )
    records["meth_TCGA_dis"] = correlate(
        gene_means, dis_methylation, "TCGA450K(gene level) vs CPTAC-dis"
    )
    records["meth_ind_dis"] = correlate(
        ind_methylation, dis_methylation, "CPTAC-ind vs CPTAC-dis"
    )

    out_path = work("protein_bridge_anchors")
    ensure_dir(out_path.parent)
    with out_path.open("w", encoding="utf-8") as handle:
        json.dump(records, handle, indent=1, ensure_ascii=False)
    print(
        "\nWrote anchor records to %s and the gene-level methylation matrix to %s"
        % (out_path, gene_level_path)
    )


if __name__ == "__main__":
    main()
